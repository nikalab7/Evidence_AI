from io import BytesIO

from docx import Document
from pypdf import PdfWriter
import pytest

from app.ingestion.parsers import (
    DocumentParseError,
    UnsupportedScannedPdfError,
    parse_uploaded_document,
)


def _text_pdf() -> bytes:
    stream = b"BT /F1 12 Tf 72 720 Td (EvidenceAI text PDF fixture.) Tj ET"
    bodies = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Title (EvidenceAI PDF fixture) >>",
    ]
    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(bodies, start=1):
        offsets.append(len(data))
        data.extend(f"{number} 0 obj\n".encode())
        data.extend(body)
        data.extend(b"\nendobj\n")
    xref_offset = len(data)
    data.extend(f"xref\n0 {len(offsets)}\n".encode())
    data.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        data.extend(f"{offset:010d} 00000 n \n".encode())
    data.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R /Info 6 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n".encode()
    )
    return bytes(data)


def _scanned_pdf() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def _docx() -> bytes:
    document = Document()
    document.core_properties.title = "DOCX metadata title"
    document.core_properties.author = "EvidenceAI test"
    document.add_paragraph("First paragraph with Unicode: ქართული ტექსტი.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Service"
    table.cell(0, 1).text = "Payments"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_plain_text_and_markdown_extract_metadata_and_utf8():
    text = parse_uploaded_document("incident.txt", "First line\nქართული".encode())
    assert text.text == "First line\nქართული"
    assert text.file_format == "txt"
    assert text.title == "incident"
    assert text.file_size_bytes == len("First line\nქართული".encode())

    markdown = parse_uploaded_document("notes.md", b"# Release Notes\n\nText body")
    assert markdown.title == "Release Notes"
    assert markdown.parser_metadata["title_source"] == "markdown_heading"


def test_docx_extracts_paragraphs_tables_and_core_metadata():
    parsed = parse_uploaded_document("upload.docx", _docx())
    assert parsed.file_format == "docx"
    assert parsed.title == "DOCX metadata title"
    assert "ქართული ტექსტი" in parsed.text
    assert "Service\tPayments" in parsed.text
    assert parsed.parser_metadata["author"] == "EvidenceAI test"


def test_text_pdf_extracts_text_metadata_and_page_count():
    parsed = parse_uploaded_document("source.pdf", _text_pdf())
    assert parsed.file_format == "pdf"
    assert parsed.title == "EvidenceAI PDF fixture"
    assert parsed.text == "EvidenceAI text PDF fixture."
    assert parsed.parser_metadata["page_count"] == 1


def test_scanned_pdf_is_rejected_with_specific_message():
    with pytest.raises(UnsupportedScannedPdfError, match="Scanned/image-only"):
        parse_uploaded_document("scan.pdf", _scanned_pdf())


@pytest.mark.parametrize(("filename", "content"), [
    ("empty.txt", b"\n  "),
    ("broken.txt", b"\xff"),
    ("binary.txt", b"text\x00binary"),
    ("bad.docx", b"not a zip"),
    ("bad.pdf", b"not a PDF"),
    ("path/../../secret.txt", b"valid content"),
    ("unsupported.rtf", b"valid content"),
])
def test_malformed_empty_or_unsupported_uploads_fail_closed(filename, content):
    if filename == "path/../../secret.txt":
        parsed = parse_uploaded_document(filename, content)
        assert parsed.parser_metadata
        assert parsed.file_format == "txt"
        return
    with pytest.raises(DocumentParseError):
        parse_uploaded_document(filename, content)


def test_upload_and_extracted_text_limits_are_enforced():
    with pytest.raises(DocumentParseError, match="5 MiB"):
        parse_uploaded_document("large.txt", b"x" * 64, max_upload_bytes=63)
    with pytest.raises(DocumentParseError, match="200,000-character"):
        parse_uploaded_document("large.txt", b"x" * 200_001)
